#!/usr/bin/env python3
"""
Rebuild the Manual Entry sheet in Test_Inspection_Plan_Rev_3.xlsm
implementing the Rev 7 card layout from the VBA specification.
Saves output as Test_Inspection_Plan_Rev_7.xlsm preserving all VBA.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

SRC = r"E:\Coding Projects\Test and Inspection Plan\Test_Inspection_Plan_Rev_3.xlsm"
OUT = r"E:\Coding Projects\Test and Inspection Plan\Test_Inspection_Plan_Rev_7.xlsm"

CARDS_START     = 6
ROWS_PER_CARD   = 13
NUM_CARDS       = 50

C_M = "2F5496"   # Medium border: RGB(47, 84, 150)
C_T = "BFBFBF"   # Thin border:   RGB(191, 191, 191)

# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------
def s(style, color="000000"):
    return Side(style=style, color=color) if style else Side(style=None)

def bdr(left=None, right=None, top=None, bottom=None):
    n = Side(style=None)
    return Border(left=left or n, right=right or n, top=top or n, bottom=bottom or n)

def solid(hex_c):
    return PatternFill(patternType="solid", fgColor=hex_c)

def fnt(name, size, bold=False, color="000000"):
    return Font(name=name, size=size, bold=bold, color=color)

def merge_row(ws, row, c1, c2, value, bg, font_obj,
              halign="left", valign="center", wrap=False):
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    # Only the top-left cell of a merged range is writable
    cell = ws.cell(row, c1)
    cell.value = value
    cell.fill = solid(bg)
    cell.font = font_obj
    cell.alignment = Alignment(horizontal=halign, vertical=valign, wrap_text=wrap)
    return cell

# ---------------------------------------------------------------------------
# Border pattern helpers (translated directly from VBA)
# ---------------------------------------------------------------------------
def outline_med(ws, row, c1, c2):
    """Medium all sides, then thin grey bottom (card header treatment)."""
    med, tg = s("medium", C_M), s("thin", C_T)
    for c in range(c1, c2 + 1):
        L = med if c == c1 else s(None)
        R = med if c == c2 else s(None)
        ws.cell(row, c).border = bdr(L, R, med, tg)

def sides_med_tb(ws, row, c1, c2, bot_med=False):
    """Left/Right medium, top thin grey, bottom medium or thin grey."""
    med, tg = s("medium", C_M), s("thin", C_T)
    bot = med if bot_med else tg
    for c in range(c1, c2 + 1):
        L = med if c == c1 else s(None)
        R = med if c == c2 else s(None)
        ws.cell(row, c).border = bdr(L, R, tg, bot)

def left_cell(ws, row, c1, c2, bot_med=False):
    """Left medium, right thin grey."""
    med, tg = s("medium", C_M), s("thin", C_T)
    bot = med if bot_med else tg
    for c in range(c1, c2 + 1):
        L = med if c == c1 else s(None)
        R = tg if c == c2 else s(None)
        ws.cell(row, c).border = bdr(L, R, tg, bot)

def mid_cell(ws, row, col, bot_med=False):
    """Thin grey all sides."""
    tg = s("thin", C_T)
    bot = s("medium", C_M) if bot_med else tg
    ws.cell(row, col).border = bdr(tg, tg, tg, bot)

def mid_range(ws, row, c1, c2, bot_med=False):
    """Like mid_cell but for a merged range (left on c1, right on c2)."""
    tg = s("thin", C_T)
    bot = s("medium", C_M) if bot_med else tg
    for c in range(c1, c2 + 1):
        L = tg if c == c1 else s(None)
        R = tg if c == c2 else s(None)
        ws.cell(row, c).border = bdr(L, R, tg, bot)

def label_row_borders(ws, row, c1, c2):
    """Left/right medium, inner vertical thin grey, top/bot thin grey."""
    med, tg = s("medium", C_M), s("thin", C_T)
    for c in range(c1, c2 + 1):
        L = med if c == c1 else tg
        R = med if c == c2 else tg
        ws.cell(row, c).border = bdr(L, R, tg, tg)

def right_border_only(ws, row, col, bot_med=False):
    """Right medium, top thin grey, bottom thin grey/medium, no left."""
    tg, med = s("thin", C_T), s("medium", C_M)
    bot = med if bot_med else tg
    ws.cell(row, col).border = bdr(s(None), med, tg, bot)

# ---------------------------------------------------------------------------
# Build a single Rev-7 manual card
# ---------------------------------------------------------------------------
def build_card(ws, start_row, card_num):
    r0  = start_row
    r1  = r0 + 1
    r2  = r0 + 2
    r3  = r0 + 3
    r4  = r0 + 4
    r5  = r0 + 5
    r6  = r0 + 6
    r7  = r0 + 7
    r12 = r0 + 12

    heights = {
        r0: 20, r1: 14, r2: 18, r3: 18, r4: 14,
        r5: 42, r6: 16, r7: 16,
        r0+8: 22, r0+9: 22, r0+10: 22, r0+11: 28, r12: 7
    }
    for rr, h in heights.items():
        ws.row_dimensions[rr].height = h

    # Reusable font objects
    f_white9b  = fnt("Arial", 9, bold=True, color="FFFFFF")
    f_white8b  = fnt("Arial", 8, bold=True, color="FFFFFF")
    f_navy8b   = fnt("Arial", 8, bold=True, color="1F3864")
    f_green8b  = fnt("Arial", 8, bold=True, color="1E4620")
    f_navy8    = fnt("Arial", 8, color="1F3864")
    f_black9   = fnt("Arial", 9, color="000000")

    # R0: Card title (dark navy, merged A:F)
    merge_row(ws, r0, 1, 6, f"  ITEM {card_num}", "1F3864", f_white9b)
    outline_med(ws, r0, 1, 6)

    # R1: Labels (amber, individual cells A-F)
    for c, lbl in enumerate(["Op No.", "Sampling", "Type", "Spec", "Units", "Bubble #"], 1):
        cell = ws.cell(r1, c)
        cell.value = lbl
        cell.fill = solid("FFD966")
        cell.font = f_navy8b
        cell.alignment = Alignment(horizontal="center", vertical="center")
    label_row_borders(ws, r1, 1, 6)

    # R2: Input cells (white, individual A-F)
    for c in range(1, 7):
        ws.cell(r2, c).fill = solid("FFFFFF")
        ws.cell(r2, c).font = f_black9
        ws.cell(r2, c).alignment = Alignment(horizontal="left", vertical="center")
    label_row_borders(ws, r2, 1, 6)

    # R3: Equipment row
    merge_row(ws, r3, 1, 2, "  Equipment Name:", "EDF7EE", f_green8b)
    left_cell(ws, r3, 1, 2, False)

    ws.cell(r3, 3).fill = solid("FFFFFF")
    ws.cell(r3, 3).font = f_black9
    ws.cell(r3, 3).alignment = Alignment(horizontal="left", vertical="center")
    mid_cell(ws, r3, 3, False)

    ws.cell(r3, 4).value = "  Equipment No.:"
    ws.cell(r3, 4).fill = solid("EDF7EE")
    ws.cell(r3, 4).font = f_green8b
    ws.cell(r3, 4).alignment = Alignment(horizontal="left", vertical="center")
    mid_cell(ws, r3, 4, False)

    ws.cell(r3, 5).fill = solid("FFFFFF")
    ws.cell(r3, 5).font = f_black9
    ws.cell(r3, 5).alignment = Alignment(horizontal="left", vertical="center")
    mid_cell(ws, r3, 5, False)

    ws.cell(r3, 6).fill = PatternFill()
    right_border_only(ws, r3, 6, False)

    # R4: Description header (dark gold, merged A:F)
    merge_row(ws, r4, 1, 6,
              "  INSPECTION REQUIREMENT / DESCRIPTION:", "A67C00", f_white8b)
    sides_med_tb(ws, r4, 1, 6, False)

    # R5: Description input (yellow, merged A:F)
    merge_row(ws, r5, 1, 6, None, "FFFDE7", f_black9,
              halign="left", valign="top", wrap=True)
    sides_med_tb(ws, r5, 1, 6, False)

    # R6: Ref / Process Ref row
    merge_row(ws, r6, 1, 3, "  Ref: ", "DEEAF6", f_navy8)
    left_cell(ws, r6, 1, 3, False)

    merge_row(ws, r6, 4, 5, "  Process Ref: ", "DEEAF6", f_navy8)
    mid_range(ws, r6, 4, 5, False)

    ws.cell(r6, 6).fill = PatternFill()
    right_border_only(ws, r6, 6, False)

    # R7: TAR header (dark green, merged A:F)
    merge_row(ws, r7, 1, 6, "  TAR HOLDER - PLEASE COMPLETE:", "1E4620", f_white8b)
    sides_med_tb(ws, r7, 1, 6, False)

    # R8-R11: TAR input rows
    tar_l = ["Verification Method", "MSA Verification Method",
             "Responsible Person", "Reaction Plan"]
    tar_r = ["Equipment Type", "MSA Ref / Justification",
             "Control Method", "Comments"]

    for idx in range(4):
        tr      = r7 + 1 + idx
        is_last = idx == 3

        merge_row(ws, tr, 1, 2, f"  {tar_l[idx]}:", "EDF7EE", f_green8b)
        left_cell(ws, tr, 1, 2, is_last)

        ws.cell(tr, 3).fill = solid("FFFFFF")
        ws.cell(tr, 3).font = f_black9
        ws.cell(tr, 3).alignment = Alignment(horizontal="left", vertical="center")
        mid_cell(ws, tr, 3, is_last)

        ws.cell(tr, 4).value = f"  {tar_r[idx]}:"
        ws.cell(tr, 4).fill = solid("EDF7EE")
        ws.cell(tr, 4).font = f_green8b
        ws.cell(tr, 4).alignment = Alignment(horizontal="left", vertical="center")
        mid_cell(ws, tr, 4, is_last)

        ws.cell(tr, 5).fill = solid("FFFFFF")
        ws.cell(tr, 5).font = f_black9
        ws.cell(tr, 5).alignment = Alignment(horizontal="left", vertical="center")
        mid_cell(ws, tr, 5, is_last)

        ws.cell(tr, 6).fill = PatternFill()
        right_border_only(ws, tr, 6, is_last)

    # R12: Spacer – clear everything
    for c in range(1, 7):
        ws.cell(r12, c).value = None
        ws.cell(r12, c).fill = PatternFill()
        ws.cell(r12, c).border = Border()
        ws.cell(r12, c).alignment = Alignment()

# ---------------------------------------------------------------------------
# Rebuild the Manual Entry sheet
# ---------------------------------------------------------------------------
def rebuild_manual_entry(ws):
    # 1. Unmerge all merged cells in the card area
    for m in list(ws.merged_cells.ranges):
        if m.min_row >= CARDS_START:
            ws.unmerge_cells(str(m))

    # 2. Clear content and formatting
    clear_to = CARDS_START + NUM_CARDS * ROWS_PER_CARD + 10
    for r in range(CARDS_START, clear_to + 1):
        for c in range(1, 7):
            cell = ws.cell(r, c)
            cell.value     = None
            cell.fill      = PatternFill()
            cell.font      = Font()
            cell.border    = Border()
            cell.alignment = Alignment()

    # 3. Clear existing data validations
    ws.data_validations.dataValidation.clear()

    # 4. Build 50 cards
    for i in range(1, NUM_CARDS + 1):
        start_row = CARDS_START + (i - 1) * ROWS_PER_CARD
        build_card(ws, start_row, i)

    # 5. Batch data validations (one rule per column, sqref covers all cards)
    vm_cells  = " ".join(
        f"C{CARDS_START + (i-1)*ROWS_PER_CARD + 8}"
        for i in range(1, NUM_CARDS + 1)
    )
    et_cells  = " ".join(
        f"E{CARDS_START + (i-1)*ROWS_PER_CARD + 8}"
        for i in range(1, NUM_CARDS + 1)
    )
    msa_cells = " ".join(
        f"C{CARDS_START + (i-1)*ROWS_PER_CARD + 9}"
        for i in range(1, NUM_CARDS + 1)
    )

    dv_vm  = DataValidation(type="list",
                            formula1="DropdownLists!$A$2:$A$8",
                            allow_blank=True, showDropDown=False)
    dv_et  = DataValidation(type="list",
                            formula1="DropdownLists!$C$2:$C$99",
                            allow_blank=True, showDropDown=False)
    dv_msa = DataValidation(type="list",
                            formula1="DropdownLists!$B$2:$B$12",
                            allow_blank=True, showDropDown=False)

    dv_vm.sqref  = vm_cells
    dv_et.sqref  = et_cells
    dv_msa.sqref = msa_cells

    ws.add_data_validation(dv_vm)
    ws.add_data_validation(dv_et)
    ws.add_data_validation(dv_msa)

    # 6. Update header subtitle to reflect Rev 7
    ws.cell(2, 1).value = "TEST & INSPECTION PLAN  -  REV 7  (Manual Entry)"

# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    print("Loading workbook ...")
    wb = openpyxl.load_workbook(SRC, keep_vba=True)

    ws_manual = wb["Manual Entry"]
    print(f"Rebuilding {NUM_CARDS} cards on 'Manual Entry' sheet ...")
    rebuild_manual_entry(ws_manual)

    # Also update the Inspection Report subtitle
    ws_report = wb["Inspection Report"]
    if ws_report.cell(2, 1).value and "REV" in str(ws_report.cell(2, 1).value):
        ws_report.cell(2, 1).value = ws_report.cell(2, 1).value.replace("REV 3", "REV 7")

    wb.save(OUT)
    print(f"Saved to: {OUT}")
