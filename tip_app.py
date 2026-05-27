#!/usr/bin/env python3
"""Test & Inspection Plan - Standalone Application  Rev 7"""

import customtkinter as ctk
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import json, os, datetime, tempfile
from pathlib import Path

try:
    import win32com.client
    OUTLOOK_OK = True
except ImportError:
    OUTLOOK_OK = False

import openpyxl
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import cm, mm

# ── Colour palette ──────────────────────────────────────────────────────────
C_NAVY    = "#1F3864"
C_NAVY2   = "#2F5496"
C_AMBER   = "#FFD966"
C_DGOLD   = "#A67C00"
C_CREAM   = "#FFFDE7"
C_SKYBLUE = "#DEEAF6"
C_LTGREEN = "#EDF7EE"
C_DKGREEN = "#1E4620"
C_GREY    = "#BFBFBF"

# ── Dropdown data ────────────────────────────────────────────────────────────
VM_OPTIONS = ["", "CMS", "MANUAL", "VISUAL", "PROCESS CONTROL PLAN",
              "SPECIAL PROCESS CONTROL", "OTHER", "N/A"]

MSA_OPTIONS = ["",
    "FULL GAUGE R&R AND BIAS STUDY", "MEASUREMENT UNCERTAINTY ANALYSIS",
    "ATTRIBUTE AGREEMENT ANALYSIS",
    "READ-ACROSS FROM AN APPROVED MEASUREMENT STUDY REPORT",
    "READ-ACROSS FROM DIMENSIONAL MEASUREMENT EQUIPMENT GUIDE MXG011",
    "TYPE 1 WITH BIAS STUDY", "AUTOMATED MEASUREMENT SYSTEM AND BIAS STUDY",
    "APPROVED CONTROL PLAN", "MSA NOT REQUIRED", "SPECIAL PROCESS CONTROL", "N/A"]

ET_OPTIONS = ["",
    "3D STRUCTURED LASER SCANNING","ACCELEROMETER","AIR GAUGE","AIRFLOW",
    "ANGLE PLATE","ARTICULATED ARM (eg. FARO / ROMOR)","AUTOMATIC HEIGHT GAUGE",
    "BINOCULARS","BORE MICROMETER  2 - POINT","BORE MICROMETER  3 - POINT",
    "BORE INDICATOR  2 - POINT","BORE INDICATOR  3 - POINT","BOROSCOPE",
    "CALIPER - DEPTH","CALIPER - DIAL","CALIPER - DIGITAL","CALIPER - VERNIER",
    "CALIPER (DIGITAL) & SCREW THREAD PLUG GAUGE","CALIPER / VISUAL",
    "CHAMFER GAUGE","CIRC TAPE","CMM",
    "CNC MACHINE  - IN CYCLE PROBING (CONTROL PLAN)",
    "COMPARATOR GAUGE","COMPARATOR GAUGE WITH DTI","COUNTERSINK GAUGE",
    "DEPTH GAUGE","DIATEST","DIGITAL HEIGHT GAUGE","DIGITAL PROTRACTOR",
    "DTI (CLOCK)","ELECTRICAL CONDUCTIVITY METER","ELECTRICAL CONTINUITY TESTER",
    "ENGINEERS SQUARE","FEELER GAUGE","GAUGE BLOCK & FEELER GAUGE","GAUGE BLOCKS",
    "GRATICULE","HAND HELD LASER SCANNER","HARDNESS TESTER","HELI-COIL DEPTH",
    "LASER TRACKER","MAGNIFICATION DEVICE","MAGNETIC PARTICLE INSPECTION",
    "MANUAL CMM","MICROMETER - ANALOGUE","MICROMETER - DIGITAL",
    "MICROMETER - INTERNAL","MICROMETER - DEEP THROAT","MICROMETER - DEPTH",
    "MICROMETER - POINTED MIC","MICROSCOPE","MODULAR BEAM COMPARATOR",
    "MOMENT WEIGHING SCALES","MOULD AND SHADOWGRAPH",
    "MOULD, SHADOWGRAPH & GRATICULE","OMNIGAUGE COMPARITOR","OTHER",
    "OPTICAL 3D MEASUREMENT DEVICE (eg ALICONA)","TOOLSETTER (eg PARLEC)",
    "PLAIN PLUG GAUGE","PRESSURE GAUGE","PRE-SWAGE INSERTS","PROFILE PROJECTOR",
    "PROTRACTOR (HELIX - GRATICULE)","PROTRACTOR (VERNIER)","RADIOGRAPHY",
    "RADIUS GAUGE","ROLLS-ROYCE APPROVED TOOLING/GAUGING","ROTARY TABLE WITH DTI",
    "ROUNDNESS & FORM MACHINES","ROUND UP TABLE","RULER",
    "SCREW THREAD PLUG GAUGES","SECTION CALIPER","SETTING BAR","SETTING RING",
    "SHADOWGRAPH","SNAP GAUGE","SPECIAL TO PRODUCT - DEPTH GAUGE",
    "SPECIAL TO PRODUCT GAUGE","SPIRIT LEVEL","SQUARE GAUGE",
    "STRAIGHT EDGE & DEPTH GAUGE",
    "SURFACE FINISH COMPARATOR (eg.RUBERT BLOCKS)","SURFACE ROUGHNESS METER",
    "TAPE MEASURE","THERMOMETER","THICKNESS METER","THREAD GAUGES",
    "THREAD PLUG GAUGES","TORQUE WRENCH","ULTRASOUND THICKNESS GAUGE",
    "UNIVERSAL BEVEL PROTRACTOR","VISUAL VERIFICATION","WEIGHING SCALES",
    "WIRE GAUGES","N/A"]

# ── Data model ───────────────────────────────────────────────────────────────
def new_card(n):
    return dict(item_num=n, op_no="", sampling="", meas_type="", spec="",
                units="", bubble="", equipment_name="", equipment_no="",
                description="", ref_loc="", process_ref="",
                verification_method="", equipment_type="", msa_method="",
                msa_ref="", responsible_person="", control_method="",
                reaction_plan="", comments="")

def new_plan():
    return dict(mva_author="", mva_number="",
                date=datetime.date.today().strftime("%d/%m/%Y"),
                mode="manual")

def new_tar_holder():
    return dict(name="", date=datetime.date.today().strftime("%d/%m/%Y"),
                employer_no="")

# ── Persistence ──────────────────────────────────────────────────────────────
def save_plan(path, plan, cards, tar_holder=None):
    data = {"version": "1.1", "plan": plan, "cards": cards}
    if tar_holder is not None:
        data["tar_holder"] = tar_holder
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)

def load_plan(path):
    with open(path, "r", encoding="utf-8") as f:
        d = json.load(f)
    return (d.get("plan", new_plan()),
            d.get("cards", []),
            d.get("tar_holder", new_tar_holder()))

# ── NetInspect import ────────────────────────────────────────────────────────
NI = dict(op=4, sampling=7, desc=8, type=9, spec=10,
          units=11, bubble=12, proc=15, ref=18)

def import_netinspect(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active
    hdr = str(ws.cell(2, NI["op"]).value or "")
    if "operation" not in hdr.lower():
        raise ValueError("Not a standard NetInspect export (expected 'Operation' in col D, row 2).")
    cards = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if len(row) < max(NI.values()):
            continue
        op = str(row[NI["op"]-1] or "").strip()
        if not op:
            continue
        c = new_card(len(cards) + 1)
        c["op_no"]       = op
        c["sampling"]    = str(row[NI["sampling"]-1] or "")
        c["description"] = str(row[NI["desc"]-1]     or "")
        c["meas_type"]   = str(row[NI["type"]-1]     or "")
        c["spec"]        = str(row[NI["spec"]-1]      or "")
        c["units"]       = str(row[NI["units"]-1]     or "")
        c["bubble"]      = str(row[NI["bubble"]-1]    or "")
        c["ref_loc"]     = str(row[NI["ref"]-1]       or "")
        c["process_ref"] = str(row[NI["proc"]-1]      or "")
        cards.append(c)
    wb.close()
    return cards

# ── PDF export ───────────────────────────────────────────────────────────────
def _rl_color(h):
    r, g, b = int(h[0:2],16), int(h[2:4],16), int(h[4:6],16)
    return colors.Color(r/255, g/255, b/255)

RL_NAVY  = _rl_color("1F3864"); RL_GREEN = _rl_color("1E4620")
RL_AMBER = _rl_color("FFD966"); RL_GOLD  = _rl_color("A67C00")
RL_CREAM = _rl_color("FFFDE7"); RL_SKY   = _rl_color("DEEAF6")
RL_LTGRN = _rl_color("EDF7EE"); RL_GREY  = _rl_color("BFBFBF")

def _ps(name, size=8, bold=False, color=colors.black):
    fn = "Helvetica-Bold" if bold else "Helvetica"
    return ParagraphStyle(name, fontName=fn, fontSize=size,
                          textColor=color, leading=size+3)

PS_WH9B = _ps("wh9b", 9, True,  colors.white)
PS_WH8B = _ps("wh8b", 8, True,  colors.white)
PS_NV8B = _ps("nv8b", 8, True,  RL_NAVY)
PS_GR8B = _ps("gr8b", 8, True,  RL_GREEN)
PS_NV8  = _ps("nv8",  8, False, RL_NAVY)
PS_BK9  = _ps("bk9",  9, False, colors.black)

def _p(txt, style): return Paragraph(str(txt) if txt else "", style)

def _grid(rows_data, col_w, bg_pairs=None, vpad=3, hpad=4):
    ts = [("GRID",(0,0),(-1,-1),0.4,RL_GREY),
          ("LEFTPADDING",(0,0),(-1,-1),hpad),
          ("RIGHTPADDING",(0,0),(-1,-1),hpad),
          ("TOPPADDING",(0,0),(-1,-1),vpad),
          ("BOTTOMPADDING",(0,0),(-1,-1),vpad)]
    for (r,c1,c2), bg in (bg_pairs or []):
        ts.append(("BACKGROUND",(c1,r),(c2,r),bg))
    return Table(rows_data, colWidths=col_w, style=TableStyle(ts))

def export_pdf(filepath, plan, cards, tar_holder=None):
    doc = SimpleDocTemplate(filepath, pagesize=A4,
                            leftMargin=1.5*cm, rightMargin=1.5*cm,
                            topMargin=1.5*cm, bottomMargin=1.5*cm)
    W = doc.width
    story = []

    story.append(Table([[_p("TEST &amp; INSPECTION PLAN  -  Rev 7", PS_WH9B)]],
        colWidths=[W], style=TableStyle([
            ("BACKGROUND",(0,0),(-1,-1),RL_NAVY),
            ("LEFTPADDING",(0,0),(-1,-1),10),
            ("TOPPADDING",(0,0),(-1,-1),8),
            ("BOTTOMPADDING",(0,0),(-1,-1),8)])))
    story.append(Spacer(1, 4*mm))

    story.append(_grid(
        [[_p("MVA Number:",  PS_NV8B), _p(plan.get("mva_number",""), PS_BK9),
          _p("Prepared By:", PS_NV8B), _p(plan.get("mva_author",""),  PS_BK9)],
         [_p("Date:",        PS_NV8B), _p(plan.get("date",""),         PS_BK9),
          _p("Input Method:",PS_NV8B), _p(plan.get("mode","").capitalize(), PS_BK9)]],
        [W*.15, W*.35, W*.15, W*.35],
        [((0,0,3),RL_SKY), ((1,0,3),RL_SKY)]))
    story.append(Spacer(1, 6*mm))

    mode = plan.get("mode","manual")
    for card in cards:
        story.extend(_card_story(card, mode, W))
        story.append(Spacer(1, 3*mm))

    story.append(Spacer(1, 5*mm))
    story.append(Table([[_p("AUTHORISATION &amp; APPROVAL", PS_WH9B)]],
        colWidths=[W], style=TableStyle([
            ("BACKGROUND",(0,0),(-1,-1),RL_NAVY),
            ("LEFTPADDING",(0,0),(-1,-1),10),
            ("TOPPADDING",(0,0),(-1,-1),6),
            ("BOTTOMPADDING",(0,0),(-1,-1),6)])))

    th = tar_holder or {}
    story.append(_grid(
        [[_p("Role",PS_NV8B), _p("Full Name",PS_NV8B),
          _p("Signature",PS_NV8B), _p("Date",PS_NV8B), _p("Emp No.",PS_NV8B)],
         [_p("Approved By (TAR Holder)",PS_BK9),
          _p(th.get("name",""), PS_BK9),
          _p("[ digital certificate attached ]", PS_NV8) if th.get("name") else "",
          _p(th.get("date",""), PS_BK9),
          _p(th.get("employer_no",""), PS_BK9)]],
        [W*.28, W*.20, W*.22, W*.15, W*.15],
        [((0,0,4),RL_SKY)], vpad=14))

    doc.build(story)

def _card_story(card, mode, W):
    items = []
    rows = []
    row_bg = []

    if mode == "automated":
        hdr_txt = (f"  ITEM {card['item_num']}"
                   + (f"   Op: {card['op_no']}" if card.get("op_no") else "")
                   + (f"   Sampling: {card['sampling']}" if card.get("sampling") else "")
                   + (f"   Spec: {card['spec']}{' '+card['units'] if card.get('units') else ''}"
                      if card.get("spec") else ""))
        rows.append([_p(hdr_txt, PS_WH9B)])
        rows.append([_p(card.get("description",""), PS_BK9)])
        row_bg = [((0,0,0),RL_NAVY), ((1,0,0),RL_CREAM)]
    else:
        rows.append([_p(f"  ITEM {card['item_num']}", PS_WH9B)])
        rows.append([_grid(
            [[_p("Op No.",PS_NV8B),  _p("Sampling",PS_NV8B), _p("Type",PS_NV8B),
              _p("Spec",PS_NV8B),    _p("Units",PS_NV8B),    _p("Bubble #",PS_NV8B)],
             [_p(card.get("op_no",""),PS_BK9),    _p(card.get("sampling",""),PS_BK9),
              _p(card.get("meas_type",""),PS_BK9), _p(card.get("spec",""),PS_BK9),
              _p(card.get("units",""),PS_BK9),     _p(card.get("bubble",""),PS_BK9)]],
            [W/6]*6, [((0,0,5),RL_AMBER)])])
        rows.append([_grid(
            [[_p("  Equipment Name:",PS_GR8B), _p(card.get("equipment_name",""),PS_BK9),
              _p("  Equipment No.:", PS_GR8B), _p(card.get("equipment_no",""),  PS_BK9)]],
            [W*.25]*4, [((0,0,3),RL_LTGRN)])])
        rows.append([_p("  INSPECTION REQUIREMENT / DESCRIPTION:", PS_WH8B)])
        rows.append([_p(card.get("description",""), PS_BK9)])
        row_bg = [((0,0,0),RL_NAVY), ((3,0,0),RL_GOLD), ((4,0,0),RL_CREAM)]

    rows.append([_grid(
        [[_p(f"  Ref: {card.get('ref_loc','')}", PS_NV8),
          _p(f"  Process Ref: {card.get('process_ref','')}", PS_NV8)]],
        [W*.5, W*.5], [((0,0,1),RL_SKY)])])

    rows.append([_p("  TAR HOLDER ASSESSMENT:", PS_WH8B)])
    row_bg.append(((len(rows)-1, 0, 0), RL_GREEN))

    rows.append([_grid(
        [[_p("  Verification Method:", PS_GR8B), _p(card.get("verification_method",""),PS_BK9),
          _p("  Equipment Type:",      PS_GR8B), _p(card.get("equipment_type",""),     PS_BK9)],
         [_p("  MSA Verification Method:",PS_GR8B), _p(card.get("msa_method",""),     PS_BK9),
          _p("  MSA Ref / Justification:",PS_GR8B), _p(card.get("msa_ref",""),         PS_BK9)],
         [_p("  Responsible Person:", PS_GR8B), _p(card.get("responsible_person",""),  PS_BK9),
          _p("  Control Method:",     PS_GR8B), _p(card.get("control_method",""),       PS_BK9)],
         [_p("  Reaction Plan:",      PS_GR8B), _p(card.get("reaction_plan",""),        PS_BK9),
          _p("  Comments:",           PS_GR8B), _p(card.get("comments",""),             PS_BK9)]],
        [W*.25]*4,
        [((r,0,0),RL_LTGRN) for r in range(4)] + [((r,2,2),RL_LTGRN) for r in range(4)])])

    ts = [("LEFTPADDING",(0,0),(-1,-1),0), ("RIGHTPADDING",(0,0),(-1,-1),0),
          ("TOPPADDING",(0,0),(-1,-1),0),  ("BOTTOMPADDING",(0,0),(-1,-1),0),
          ("BOX",(0,0),(-1,-1),1.5,RL_NAVY)]
    for (ri,c1,c2), bg in row_bg:
        ts.append(("BACKGROUND",(c1,ri),(c2,ri),bg))

    items.append(Table(rows, colWidths=[W], style=TableStyle(ts)))
    return items

# ── Email ────────────────────────────────────────────────────────────────────
def send_outlook(to_addr, subject, body, attachments=None):
    if not OUTLOOK_OK:
        raise RuntimeError("pywin32 not installed")
    ol = win32com.client.Dispatch("Outlook.Application")
    m  = ol.CreateItem(0)
    m.To, m.Subject, m.Body = to_addr, subject, body
    for a in (attachments or []):
        if os.path.exists(a):
            m.Attachments.Add(a)
    m.Display()

# ── Card dialog (Author only — no TAR fields) ────────────────────────────────
class CardDialog(ctk.CTkToplevel):
    def __init__(self, parent, card, mode, on_save):
        super().__init__(parent)
        self.card    = dict(card)
        self.mode    = mode
        self.on_save = on_save
        self.title(f"ITEM {card['item_num']}")
        self.geometry("700x520")
        self.resizable(True, True)
        self.grab_set()
        self._vars = {}
        self._build()

    def _build(self):
        scr = ctk.CTkScrollableFrame(self)
        scr.pack(fill="both", expand=True, padx=6, pady=6)
        scr.columnconfigure([1, 3], weight=1)

        hf = ctk.CTkFrame(scr, fg_color=C_NAVY, corner_radius=4)
        hf.grid(row=0, column=0, columnspan=4, sticky="ew", pady=(0, 8))
        ctk.CTkLabel(hf, text=f"  ITEM {self.card['item_num']}",
                     text_color="white",
                     font=ctk.CTkFont("Arial", 11, "bold")).pack(side="left", pady=5)

        r = 1
        if self.mode == "manual":
            for lbl, key, col_offset in [
                ("Op No.:",   "op_no",    0), ("Sampling:", "sampling", 2),
                ("Type:",     "meas_type",0), ("Spec:",     "spec",     2),
                ("Units:",    "units",    0), ("Bubble #:", "bubble",   2),
            ]:
                if col_offset == 0:
                    ro = r
                ctk.CTkLabel(scr, text=lbl, anchor="e").grid(
                    row=ro, column=col_offset, padx=(4,2), pady=2, sticky="e")
                v = ctk.StringVar(value=self.card.get(key,""))
                self._vars[key] = v
                ctk.CTkEntry(scr, textvariable=v).grid(
                    row=ro, column=col_offset+1, padx=(0,8), pady=2, sticky="ew")
                if col_offset == 2:
                    r += 1

        ef = ctk.CTkFrame(scr, fg_color=C_LTGREEN, corner_radius=0)
        ef.grid(row=r, column=0, columnspan=4, sticky="ew", pady=(8,1))
        ef.columnconfigure([1,3], weight=1)
        ctk.CTkLabel(ef, text="Equipment Name:", anchor="e").grid(
            row=0, column=0, padx=(8,2), pady=4, sticky="e")
        v = ctk.StringVar(value=self.card.get("equipment_name",""))
        self._vars["equipment_name"] = v
        ctk.CTkEntry(ef, textvariable=v).grid(
            row=0, column=1, padx=(0,8), pady=4, sticky="ew")
        ctk.CTkLabel(ef, text="Equipment No.:", anchor="e").grid(
            row=0, column=2, padx=(8,2), pady=4, sticky="e")
        v2 = ctk.StringVar(value=self.card.get("equipment_no",""))
        self._vars["equipment_no"] = v2
        ctk.CTkEntry(ef, textvariable=v2, width=140).grid(
            row=0, column=3, padx=(0,8), pady=4, sticky="ew")
        r += 1

        dh = ctk.CTkFrame(scr, fg_color=C_DGOLD, corner_radius=0)
        dh.grid(row=r, column=0, columnspan=4, sticky="ew", pady=(6,0))
        ctk.CTkLabel(dh, text="  INSPECTION REQUIREMENT / DESCRIPTION:",
                     text_color="white",
                     font=ctk.CTkFont("Arial", 8, "bold")).pack(side="left", pady=3)
        r += 1
        self._desc = ctk.CTkTextbox(scr, height=90, fg_color=C_CREAM,
                                     border_width=1, border_color=C_GREY)
        self._desc.grid(row=r, column=0, columnspan=4, sticky="ew", pady=(0,6))
        self._desc.insert("1.0", self.card.get("description",""))
        r += 1

        rf = ctk.CTkFrame(scr, fg_color=C_SKYBLUE, corner_radius=0)
        rf.grid(row=r, column=0, columnspan=4, sticky="ew", pady=1)
        rf.columnconfigure([1,3], weight=1)
        ctk.CTkLabel(rf, text="Ref Location:", anchor="e").grid(
            row=0, column=0, padx=(8,2), pady=4, sticky="e")
        vr = ctk.StringVar(value=self.card.get("ref_loc",""))
        self._vars["ref_loc"] = vr
        ctk.CTkEntry(rf, textvariable=vr).grid(
            row=0, column=1, padx=(0,8), pady=4, sticky="ew")
        ctk.CTkLabel(rf, text="Process Ref:", anchor="e").grid(
            row=0, column=2, padx=(8,2), pady=4, sticky="e")
        vp = ctk.StringVar(value=self.card.get("process_ref",""))
        self._vars["process_ref"] = vp
        ctk.CTkEntry(rf, textvariable=vp).grid(
            row=0, column=3, padx=(0,8), pady=4, sticky="ew")
        r += 1

        bf = ctk.CTkFrame(scr, fg_color="transparent")
        bf.grid(row=r, column=0, columnspan=4, pady=14)
        ctk.CTkButton(bf, text="Save", width=120, fg_color=C_NAVY,
                      command=self._save).pack(side="left", padx=10)
        ctk.CTkButton(bf, text="Cancel", width=100, fg_color="#666666",
                      command=self.destroy).pack(side="left", padx=4)

    def _save(self):
        for key, var in self._vars.items():
            self.card[key] = var.get()
        self.card["description"] = self._desc.get("1.0", "end-1c")
        self.on_save(self.card)
        self.destroy()


# ── Role selection dialog ─────────────────────────────────────────────────────
class RoleDialog(ctk.CTkToplevel):
    def __init__(self, parent):
        super().__init__(parent)
        self.result = None
        self.title("Open Plan As...")
        self.geometry("440x220")
        self.resizable(False, False)
        self.grab_set()

        ctk.CTkLabel(self, text="How are you opening this plan?",
                     font=ctk.CTkFont("Arial", 13, "bold"),
                     text_color=C_NAVY).pack(pady=(28, 6))
        ctk.CTkLabel(self, text="Choose your role to see the right interface.",
                     font=ctk.CTkFont("Arial", 10),
                     text_color="#555555").pack(pady=(0, 20))

        bf = ctk.CTkFrame(self, fg_color="transparent")
        bf.pack()
        ctk.CTkButton(bf, text="Author\n(Create / Edit Plan)", width=170, height=52,
                      fg_color=C_NAVY,
                      command=lambda: self._pick("author")).pack(side="left", padx=12)
        ctk.CTkButton(bf, text="TAR Holder\n(Review & Approve)", width=170, height=52,
                      fg_color=C_DKGREEN,
                      command=lambda: self._pick("tar")).pack(side="left", padx=12)

    def _pick(self, role):
        self.result = role
        self.destroy()


# ── Main application shell ────────────────────────────────────────────────────
class TIPApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        ctk.set_appearance_mode("light")
        ctk.set_default_color_theme("blue")
        self.title("Test & Inspection Plan  |  Rev 7")
        self.geometry("1120x760")
        self.minsize(900, 620)

        self._current_view = None
        self._build_header()
        self._build_status()
        self._content = ctk.CTkFrame(self, fg_color="transparent")
        self._content.pack(fill="both", expand=True)

        self.show_launch()

    def _build_header(self):
        h = ctk.CTkFrame(self, fg_color=C_NAVY, height=52, corner_radius=0)
        h.pack(fill="x")
        h.pack_propagate(False)
        ctk.CTkButton(h, text="Home", width=64, height=34,
                      fg_color="transparent", hover_color=C_NAVY2,
                      text_color="white", font=ctk.CTkFont("Arial", 10),
                      command=self._go_home).pack(side="left", padx=(8,0))
        ctk.CTkFrame(h, fg_color=C_NAVY2, width=1, height=30).pack(
            side="left", padx=8)
        ctk.CTkLabel(h, text="TEST & INSPECTION PLAN",
                     text_color="white",
                     font=ctk.CTkFont("Arial", 16, "bold")).pack(side="left", padx=4)
        ctk.CTkLabel(h, text="Rev 7",
                     text_color=C_AMBER,
                     font=ctk.CTkFont("Arial", 12)).pack(side="left", padx=6)
        self._mode_badge = ctk.CTkLabel(h, text="",
                                         text_color=C_AMBER,
                                         font=ctk.CTkFont("Arial", 11, "bold"))
        self._mode_badge.pack(side="right", padx=16)

    def _build_status(self):
        self._sv = tk.StringVar(value="Ready")
        sb = ctk.CTkFrame(self, fg_color="#E0E0E0", height=24, corner_radius=0)
        sb.pack(fill="x", side="bottom")
        sb.pack_propagate(False)
        ctk.CTkLabel(sb, textvariable=self._sv,
                     font=ctk.CTkFont("Arial", 10),
                     text_color="#333333").pack(side="left", padx=10)

    def status(self, msg): self._sv.set(msg)
    def set_badge(self, text): self._mode_badge.configure(text=text)

    def show_view(self, view):
        if self._current_view:
            self._current_view.destroy()
        self._current_view = view
        view.pack(fill="both", expand=True)

    def show_launch(self):
        self.set_badge("")
        self.show_view(LaunchView(self._content, self))

    def _go_home(self):
        if isinstance(self._current_view, LaunchView):
            return
        if messagebox.askyesno("Go Home",
                "Return to the home screen?\n\nAny unsaved changes will be lost."):
            self.show_launch()
            self.status("Ready")

    def show_author(self, plan=None, cards=None, tip_path=None):
        self.set_badge("[ AUTHOR MODE ]")
        self.show_view(AuthorView(self._content, self, plan, cards, tip_path))

    def show_tar_holder(self, plan, cards, tar_holder, tip_path):
        self.set_badge("[ TAR HOLDER MODE ]")
        self.show_view(TARHolderView(self._content, self,
                                     plan, cards, tar_holder, tip_path))

    def open_plan(self):
        path = filedialog.askopenfilename(
            filetypes=[("TIP Plan","*.tip"), ("All Files","*.*")],
            title="Open Test & Inspection Plan")
        if not path:
            return
        plan, cards, tar_holder = load_plan(path)

        dlg = RoleDialog(self)
        self.wait_window(dlg)

        if dlg.result == "author":
            self.show_author(plan, cards, path)
            self.status(f"Opened: {Path(path).name}  ({len(cards)} cards)")
        elif dlg.result == "tar":
            self.show_tar_holder(plan, cards, tar_holder, path)
            self.status(f"Opened for TAR review: {Path(path).name}  ({len(cards)} cards)")


# ── Launch view ───────────────────────────────────────────────────────────────
class LaunchView(ctk.CTkFrame):
    def __init__(self, parent, app):
        super().__init__(parent, fg_color="transparent")
        self.app = app
        self._build()

    def _build(self):
        centre = ctk.CTkFrame(self, fg_color="transparent")
        centre.place(relx=0.5, rely=0.46, anchor="center")

        ctk.CTkLabel(centre, text="TEST & INSPECTION PLAN",
                     font=ctk.CTkFont("Arial", 26, "bold"),
                     text_color=C_NAVY).pack(pady=(0, 4))
        ctk.CTkLabel(centre, text="Rev 7",
                     font=ctk.CTkFont("Arial", 14),
                     text_color=C_DGOLD).pack(pady=(0, 48))

        ctk.CTkButton(centre,
                      text="  Create New Test & Inspection Plan",
                      anchor="w",
                      width=360, height=60,
                      fg_color=C_NAVY,
                      font=ctk.CTkFont("Arial", 13, "bold"),
                      command=lambda: self.app.show_author()).pack(pady=10)

        ctk.CTkButton(centre,
                      text="  Open Existing Plan  (.tip)",
                      anchor="w",
                      width=360, height=60,
                      fg_color=C_NAVY2,
                      font=ctk.CTkFont("Arial", 13, "bold"),
                      command=self.app.open_plan).pack(pady=10)


# ── Author view (3-step) ──────────────────────────────────────────────────────
class AuthorView(ctk.CTkFrame):
    def __init__(self, parent, app, plan=None, cards=None, tip_path=None):
        super().__init__(parent, fg_color="transparent")
        self.app       = app
        self.plan      = plan or new_plan()
        self.cards     = list(cards or [])
        self._tip_path = tip_path
        self._step     = 1
        self._vars     = {}
        self._mode_var = ctk.StringVar(value=self.plan.get("mode","manual"))
        self._build()

    # ── Shell ────────────────────────────────────────────────────────────────
    def _build(self):
        sidebar = ctk.CTkFrame(self, fg_color=C_NAVY, width=182, corner_radius=0)
        sidebar.pack(side="left", fill="y")
        sidebar.pack_propagate(False)

        ctk.CTkLabel(sidebar, text="PLAN STEPS",
                     text_color=C_AMBER,
                     font=ctk.CTkFont("Arial", 10, "bold")).pack(
            pady=(22,14), padx=14, anchor="w")

        self._step_btns = []
        for num, label in [("1","Plan Details"),
                            ("2","Inspection Cards"),
                            ("3","Send to TAR Holder")]:
            btn = ctk.CTkButton(
                sidebar, text=f"  {num}.  {label}",
                anchor="w", width=164, height=44,
                fg_color="transparent", hover_color=C_NAVY2,
                text_color="white", font=ctk.CTkFont("Arial", 11),
                command=lambda n=int(num): self._goto(n))
            btn.pack(padx=9, pady=2)
            self._step_btns.append(btn)

        ctk.CTkFrame(sidebar, fg_color=C_NAVY2, height=1).pack(
            fill="x", padx=10, pady=18)

        ctk.CTkButton(sidebar, text="Save Plan", width=164, height=36,
                      fg_color="#3a3a3a", hover_color="#505050",
                      command=self._save_plan).pack(padx=9, pady=3)
        ctk.CTkButton(sidebar, text="Load Plan", width=164, height=36,
                      fg_color="#3a3a3a", hover_color="#505050",
                      command=self._load_plan).pack(padx=9, pady=3)

        self._area = ctk.CTkFrame(self, fg_color="transparent")
        self._area.pack(side="left", fill="both", expand=True)

        self._frame = None
        self._goto(1)

    def _goto(self, n):
        self._flush_vars()
        self._step = n
        for i, btn in enumerate(self._step_btns, 1):
            btn.configure(fg_color=C_NAVY2 if i == n else "transparent")
        if self._frame:
            self._frame.destroy()
        self._frame = ctk.CTkFrame(self._area, fg_color="transparent")
        self._frame.pack(fill="both", expand=True, padx=14, pady=10)
        [self._step1, self._step2, self._step3][n-1](self._frame)

    def _step_header(self, parent, num, title):
        hf = ctk.CTkFrame(parent, fg_color=C_SKYBLUE, height=48, corner_radius=4)
        hf.pack(fill="x", pady=(0,14))
        hf.pack_propagate(False)
        ctk.CTkLabel(hf, text=f"  Step {num}",
                     font=ctk.CTkFont("Arial", 14, "bold"),
                     text_color=C_NAVY).pack(side="left", padx=(14,4))
        ctk.CTkLabel(hf, text=f"—  {title}",
                     font=ctk.CTkFont("Arial", 13),
                     text_color=C_NAVY).pack(side="left")

    def _flush_vars(self):
        for k, v in self._vars.items():
            self.plan[k] = v.get()
        self.plan["mode"] = self._mode_var.get()

    # ── Step 1 ───────────────────────────────────────────────────────────────
    def _step1(self, parent):
        self._step_header(parent, 1, "Plan Details")
        frm = ctk.CTkScrollableFrame(parent)
        frm.pack(fill="both", expand=True)
        frm.columnconfigure(1, weight=1)

        def row(label, key, r, today=False):
            ctk.CTkLabel(frm, text=label, anchor="e", width=210).grid(
                row=r, column=0, padx=(4,6), pady=8, sticky="e")
            v = ctk.StringVar(value=self.plan.get(key,""))
            self._vars[key] = v
            ctk.CTkEntry(frm, textvariable=v, width=360).grid(
                row=r, column=1, sticky="ew", padx=(0,4), pady=8)
            if today:
                ctk.CTkButton(frm, text="Today", width=70,
                              command=lambda vv=v: vv.set(
                                  datetime.date.today().strftime("%d/%m/%Y"))
                              ).grid(row=r, column=2, padx=4, pady=8)

        row("MVA Author (Prepared By):", "mva_author", 0)
        row("MVA Number:",               "mva_number", 1)
        row("Date (DD/MM/YYYY):",        "date",        2, today=True)

        ctk.CTkLabel(frm, text="Input Method:", anchor="e", width=210).grid(
            row=3, column=0, padx=(4,6), pady=8, sticky="e")
        mf = ctk.CTkFrame(frm, fg_color="transparent")
        mf.grid(row=3, column=1, sticky="w", pady=8)
        ctk.CTkRadioButton(mf, text="Manual",
                           variable=self._mode_var, value="manual").pack(
            side="left", padx=14)
        ctk.CTkRadioButton(mf, text="Automated (import NetInspect)",
                           variable=self._mode_var, value="automated").pack(
            side="left", padx=14)

        ctk.CTkFrame(frm, fg_color=C_GREY, height=1).grid(
            row=4, column=0, columnspan=3, sticky="ew", padx=4, pady=16)

        bf = ctk.CTkFrame(frm, fg_color="transparent")
        bf.grid(row=5, column=0, columnspan=3, pady=4)
        ctk.CTkButton(bf, text="Next: Inspection Cards  →",
                      width=230, height=40, fg_color=C_NAVY,
                      command=lambda: self._goto(2)).pack()

    # ── Step 2 ───────────────────────────────────────────────────────────────
    def _step2(self, parent):
        self._step_header(parent, 2, "Inspection Cards")
        mode = self.plan.get("mode","manual")

        tb = ctk.CTkFrame(parent, fg_color="#E8EEF7", height=44, corner_radius=4)
        tb.pack(fill="x", pady=(0,3))
        tb.pack_propagate(False)

        self._cnt = ctk.CTkLabel(tb, text="", font=ctk.CTkFont("Arial",11),
                                  text_color=C_NAVY)
        self._cnt.pack(side="right", padx=14)

        if mode == "automated":
            ctk.CTkButton(tb, text="Import NetInspect File", width=200,
                          fg_color=C_NAVY,
                          command=self._import_ni).pack(side="left", padx=8, pady=6)
        else:
            for txt, cmd, clr in [
                ("+ Add Card",      self._add_card,    C_NAVY),
                ("Edit Selected",   self._edit_card,   C_NAVY2),
                ("Delete Selected", self._del_card,    "#8B0000"),
                ("Clear All",       self._clear_cards, "#666666"),
            ]:
                ctk.CTkButton(tb, text=txt, width=130, fg_color=clr,
                              command=cmd).pack(side="left", padx=4, pady=6)

        tf = ctk.CTkFrame(parent, fg_color="white")
        tf.pack(fill="both", expand=True, pady=(2,4))
        cols = ("#","Op No.","Sampling","Type","Spec / Units","Description")
        self._tree = ttk.Treeview(tf, columns=cols, show="headings",
                                   selectmode="browse")
        cw = [38,72,80,80,130,0]
        for col, w in zip(cols, cw):
            self._tree.heading(col, text=col)
            self._tree.column(col, width=w, stretch=(w==0))
        vsb = ttk.Scrollbar(tf, orient="vertical", command=self._tree.yview)
        self._tree.configure(yscrollcommand=vsb.set)
        self._tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")
        self._tree.bind("<Double-1>", lambda e: self._edit_card())
        s = ttk.Style()
        s.configure("Treeview", font=("Arial",9), rowheight=22)
        s.configure("Treeview.Heading", font=("Arial",9,"bold"))
        self._rebuild_tree()

        bf = ctk.CTkFrame(parent, fg_color="transparent")
        bf.pack(pady=6)
        ctk.CTkButton(bf, text="← Back: Plan Details", width=180, height=36,
                      fg_color="#666666",
                      command=lambda: self._goto(1)).pack(side="left", padx=8)
        ctk.CTkButton(bf, text="Next: Send to TAR Holder  →", width=230, height=36,
                      fg_color=C_NAVY,
                      command=lambda: self._goto(3)).pack(side="left", padx=8)

    def _rebuild_tree(self):
        self._tree.delete(*self._tree.get_children())
        for c in self.cards:
            self._tree.insert("","end", values=(
                c["item_num"], c.get("op_no",""), c.get("sampling",""),
                c.get("meas_type",""),
                f"{c.get('spec','')} {c.get('units','')}".strip(),
                (c.get("description") or "")[:70]))
        n = len(self.cards)
        if hasattr(self, "_cnt"):
            self._cnt.configure(text=f"{n} card{'s' if n!=1 else ''}")

    def _import_ni(self):
        path = filedialog.askopenfilename(
            filetypes=[("Excel","*.xlsx *.xls *.xlsm"),("All Files","*.*")],
            title="Select NetInspect Export File")
        if not path: return
        try:
            self.cards = import_netinspect(path)
            self._rebuild_tree()
            self.app.status(f"Imported {len(self.cards)} items from {Path(path).name}")
        except Exception as ex:
            messagebox.showerror("Import Error", str(ex))

    def _add_card(self):
        card = new_card(len(self.cards)+1)
        def on_save(c):
            self.cards.append(c)
            self._rebuild_tree()
            self.app.status(f"Card {c['item_num']} added.")
        CardDialog(self, card, self.plan.get("mode","manual"), on_save)

    def _edit_card(self):
        sel = self._tree.selection()
        if not sel:
            messagebox.showinfo("Edit","Select a card first."); return
        idx = self._tree.index(sel[0])
        def on_save(c):
            self.cards[idx] = c
            self._rebuild_tree()
            self.app.status(f"Card {c['item_num']} updated.")
        CardDialog(self, self.cards[idx], self.plan.get("mode","manual"), on_save)

    def _del_card(self):
        sel = self._tree.selection()
        if not sel: return
        idx = self._tree.index(sel[0])
        if messagebox.askyesno("Delete",f"Delete ITEM {self.cards[idx]['item_num']}?"):
            del self.cards[idx]
            for i, c in enumerate(self.cards, 1):
                c["item_num"] = i
            self._rebuild_tree()

    def _clear_cards(self):
        if self.cards and messagebox.askyesno("Clear All","Remove all cards?"):
            self.cards = []
            self._rebuild_tree()

    # ── Step 3 ───────────────────────────────────────────────────────────────
    def _step3(self, parent):
        self._step_header(parent, 3, "Send to TAR Holder")
        frm = ctk.CTkScrollableFrame(parent)
        frm.pack(fill="both", expand=True)
        frm.columnconfigure(1, weight=1)

        p = self.plan
        summary = (f"MVA Number:   {p.get('mva_number','(not set)')}\n"
                   f"Prepared By:  {p.get('mva_author','(not set)')}\n"
                   f"Date:         {p.get('date','(not set)')}\n"
                   f"Input Method: {p.get('mode','manual').capitalize()}\n"
                   f"Cards:        {len(self.cards)}")

        ctk.CTkLabel(frm, text="PLAN SUMMARY",
                     font=ctk.CTkFont("Arial",11,"bold"),
                     text_color=C_NAVY).grid(
            row=0, column=0, columnspan=3, sticky="w", padx=8, pady=(8,2))
        ctk.CTkLabel(frm, text=summary, anchor="w", justify="left",
                     font=ctk.CTkFont("Arial",10)).grid(
            row=1, column=0, columnspan=3, sticky="w", padx=24, pady=(0,14))

        ctk.CTkLabel(frm, text="TAR Holder Email:", anchor="e", width=180).grid(
            row=2, column=0, padx=(4,6), pady=8, sticky="e")
        self._email_var = ctk.StringVar(value=self.plan.get("tar_email",""))
        ctk.CTkEntry(frm, textvariable=self._email_var, width=360).grid(
            row=2, column=1, sticky="ew", padx=(0,4), pady=8)

        ctk.CTkFrame(frm, fg_color=C_GREY, height=1).grid(
            row=3, column=0, columnspan=3, sticky="ew", padx=4, pady=14)

        bf = ctk.CTkFrame(frm, fg_color="transparent")
        bf.grid(row=4, column=0, columnspan=3, pady=4)
        ctk.CTkButton(bf, text="Send Email via Outlook", width=200,
                      fg_color=C_NAVY,
                      command=self._send_email).pack(side="left", padx=8)
        ctk.CTkButton(bf, text="Export Draft PDF", width=160,
                      fg_color=C_NAVY2,
                      command=self._export_draft).pack(side="left", padx=8)
        ctk.CTkButton(bf, text="Save Plan (.tip)", width=150,
                      fg_color="#666666",
                      command=self._save_plan).pack(side="left", padx=8)

        ctk.CTkFrame(frm, fg_color=C_GREY, height=1).grid(
            row=5, column=0, columnspan=3, sticky="ew", padx=4, pady=14)
        ctk.CTkButton(frm, text="← Back: Inspection Cards", width=190, height=36,
                      fg_color="#666666",
                      command=lambda: self._goto(2)).grid(
            row=6, column=0, columnspan=3, pady=4)

    # ── Actions ──────────────────────────────────────────────────────────────
    def _save_plan(self):
        self._flush_vars()
        path = self._tip_path or filedialog.asksaveasfilename(
            defaultextension=".tip",
            filetypes=[("TIP Plan","*.tip"),("All Files","*.*")],
            title="Save Plan")
        if not path: return
        self._tip_path = path
        save_plan(path, self.plan, self.cards)
        self.app.status(f"Saved: {path}")

    def _load_plan(self):
        path = filedialog.askopenfilename(
            filetypes=[("TIP Plan","*.tip"),("All Files","*.*")],
            title="Load Plan")
        if not path: return
        plan, cards, _ = load_plan(path)
        self.plan  = plan
        self.cards = cards
        self._tip_path = path
        self._vars.clear()
        self._mode_var.set(plan.get("mode","manual"))
        self._goto(self._step)
        self.app.status(f"Loaded: {path}  ({len(cards)} cards)")

    def _send_email(self):
        self._flush_vars()
        to = self._email_var.get().strip()
        if not to:
            messagebox.showwarning("Send Email","Enter the TAR Holder email."); return
        if not self.plan.get("mva_number"):
            messagebox.showwarning("Send Email","Enter the MVA Number first."); return
        if not self.cards:
            messagebox.showwarning("Send Email","No inspection cards to send."); return

        self.plan["tar_email"] = to
        tip_path = self._tip_path or os.path.join(
            tempfile.gettempdir(), f"TIP_{self.plan['mva_number']}.tip")
        save_plan(tip_path, self.plan, self.cards)
        self._tip_path = tip_path

        pdf_path = os.path.join(tempfile.gettempdir(),
                                f"TIP_{self.plan['mva_number']}_Draft.pdf")
        try: export_pdf(pdf_path, self.plan, self.cards)
        except Exception: pdf_path = None

        p = self.plan
        body = (
            f"Dear TAR Holder,\n\n"
            f"Please find attached the Test & Inspection Plan for your review and approval.\n\n"
            f"  MVA Number  : {p.get('mva_number','')}\n"
            f"  Date        : {p.get('date','')}\n"
            f"  Prepared By : {p.get('mva_author','')}\n\n"
            + ("All characteristics were imported from NetInspect.\n"
               if p.get("mode")=="automated" else
               "This plan was prepared using Manual Entry.\n")
            + "\nACTION REQUIRED:\n"
              "  1. Open the attached .tip file in the TIP Application\n"
              "  2. Click 'Open Existing Plan' on the home screen\n"
              "  3. Select 'TAR Holder' when prompted\n"
              "  4. Enter your details, complete all green TAR fields\n"
              "  5. Export the Approved PDF, apply your digital certificate signature\n"
              "  6. Email the signed PDF back to the author\n\n"
            f"Kind regards,\n{p.get('mva_author','[Operator]')}"
        )
        subject = f"ACTION REQUIRED: Test & Inspection Plan - {p.get('mva_number','')}"
        attachments = [a for a in [tip_path, pdf_path] if a and os.path.exists(a)]

        if OUTLOOK_OK:
            try:
                send_outlook(to, subject, body, attachments)
                self.app.status("Email draft opened in Outlook.")
            except Exception as ex:
                messagebox.showerror("Email Error", str(ex))
        else:
            messagebox.showinfo("Email (Outlook not available)",
                f"To: {to}\nSubject: {subject}\n\nPlan saved to:\n{tip_path}")

    def _export_draft(self):
        self._flush_vars()
        if not self.cards:
            messagebox.showwarning("Export PDF","No cards to export."); return
        path = filedialog.asksaveasfilename(
            defaultextension=".pdf",
            filetypes=[("PDF","*.pdf"),("All Files","*.*")],
            initialfile=f"TIP_{self.plan.get('mva_number','PLAN')}_Draft.pdf",
            title="Save Draft PDF")
        if not path: return
        try:
            export_pdf(path, self.plan, self.cards)
            self.app.status(f"Draft PDF saved: {path}")
            if messagebox.askyesno("PDF Saved","Open the PDF now?"):
                os.startfile(path)
        except Exception as ex:
            messagebox.showerror("PDF Error", str(ex))


# ── TAR Holder view ───────────────────────────────────────────────────────────
class TARHolderView(ctk.CTkFrame):
    def __init__(self, parent, app, plan, cards, tar_holder, tip_path):
        super().__init__(parent, fg_color="transparent")
        self.app        = app
        self.plan       = dict(plan)
        self.cards      = [dict(c) for c in cards]
        self.tar_holder = dict(tar_holder) if tar_holder else new_tar_holder()
        self._tip_path  = tip_path
        self._cur       = None
        self._tar_vars  = {}
        self._h_vars    = {}
        self._build()

    def _build(self):
        # ── Instruction banner ───────────────────────────────────────────────
        banner = ctk.CTkFrame(self, fg_color=C_DKGREEN, corner_radius=4)
        banner.pack(fill="x", padx=8, pady=(8,4))
        ctk.CTkLabel(
            banner,
            text=("  TAR HOLDER  —  Enter your details below, then select each inspection "
                  "item and complete the green assessment fields. Export the Approved PDF "
                  "when all items are done."),
            text_color="white",
            font=ctk.CTkFont("Arial", 10, "bold"),
            anchor="w", wraplength=950, justify="left"
        ).pack(side="left", pady=10, padx=6)

        # ── Plan info (read-only) ────────────────────────────────────────────
        info = ctk.CTkFrame(self, fg_color=C_SKYBLUE, corner_radius=4)
        info.pack(fill="x", padx=8, pady=(0,4))
        p = self.plan
        ctk.CTkLabel(info,
                     text=(f"  MVA: {p.get('mva_number','')}   |   "
                           f"Prepared By: {p.get('mva_author','')}   |   "
                           f"Date: {p.get('date','')}   |   "
                           f"{len(self.cards)} inspection items"),
                     text_color=C_NAVY,
                     font=ctk.CTkFont("Arial", 10),
                     anchor="w").pack(side="left", pady=6, padx=6)

        # ── TAR Holder details ───────────────────────────────────────────────
        det = ctk.CTkFrame(self, fg_color="#F0F5FF", corner_radius=4)
        det.pack(fill="x", padx=8, pady=(0,4))

        ctk.CTkLabel(det, text="YOUR DETAILS",
                     font=ctk.CTkFont("Arial", 10, "bold"),
                     text_color=C_NAVY).pack(anchor="w", padx=12, pady=(8,4))

        row_f = ctk.CTkFrame(det, fg_color="transparent")
        row_f.pack(fill="x", padx=12, pady=(0,10))

        def _hfield(parent, label, key, width=200):
            ctk.CTkLabel(parent, text=label,
                         font=ctk.CTkFont("Arial", 10)).pack(side="left", padx=(0,4))
            v = ctk.StringVar(value=self.tar_holder.get(key,""))
            if key == "date" and not v.get():
                v.set(datetime.date.today().strftime("%d/%m/%Y"))
            self._h_vars[key] = v
            ctk.CTkEntry(parent, textvariable=v, width=width).pack(side="left", padx=(0,18))

        _hfield(row_f, "TAR Holder Name:", "name", 220)

        # date + Today button grouped
        date_grp = ctk.CTkFrame(row_f, fg_color="transparent")
        date_grp.pack(side="left", padx=(0,18))
        ctk.CTkLabel(date_grp, text="Date (DD/MM/YYYY):",
                     font=ctk.CTkFont("Arial", 10)).pack(side="left", padx=(0,4))
        dv = ctk.StringVar(value=self.tar_holder.get("date","") or
                           datetime.date.today().strftime("%d/%m/%Y"))
        self._h_vars["date"] = dv
        ctk.CTkEntry(date_grp, textvariable=dv, width=120).pack(side="left", padx=(0,4))
        ctk.CTkButton(date_grp, text="Today", width=62, height=28,
                      command=lambda: dv.set(
                          datetime.date.today().strftime("%d/%m/%Y"))
                      ).pack(side="left")

        _hfield(row_f, "Employer Number:", "employer_no", 160)

        # ── Card treeview ────────────────────────────────────────────────────
        tf = ctk.CTkFrame(self, fg_color="white")
        tf.pack(fill="both", expand=True, padx=8, pady=(0,4))

        cols = ("#","Op No.","Description","Verif. Method","Equip. Type","Responsible")
        self._tree = ttk.Treeview(tf, columns=cols, show="headings",
                                   selectmode="browse", height=9)
        cw = [38, 72, 0, 150, 120, 130]
        for col, w in zip(cols, cw):
            self._tree.heading(col, text=col)
            self._tree.column(col, width=w, stretch=(w==0))
        vsb = ttk.Scrollbar(tf, orient="vertical", command=self._tree.yview)
        self._tree.configure(yscrollcommand=vsb.set)
        self._tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")
        self._tree.bind("<<TreeviewSelect>>", self._on_select)
        self._populate_tree()

        # ── TAR edit panel ───────────────────────────────────────────────────
        self._panel = ctk.CTkFrame(self, fg_color=C_LTGREEN, height=185)
        self._panel.pack(fill="x", padx=8, pady=(0,4))
        self._panel.pack_propagate(False)
        ctk.CTkLabel(self._panel,
                     text="Select an inspection item above to complete its TAR fields",
                     text_color="#666666",
                     font=ctk.CTkFont("Arial", 10)).pack(expand=True)

        # ── Action buttons ───────────────────────────────────────────────────
        btns = ctk.CTkFrame(self, fg_color="#E8EEF7", height=50, corner_radius=4)
        btns.pack(fill="x", padx=8, pady=(0,8))
        btns.pack_propagate(False)

        ctk.CTkButton(btns, text="Save Progress", width=160,
                      fg_color="#666666",
                      command=self._save_progress).pack(side="left", padx=8, pady=8)
        ctk.CTkButton(btns, text="Export Approved PDF", width=190,
                      fg_color=C_DKGREEN,
                      command=self._export_approved).pack(side="left", padx=4, pady=8)
        ctk.CTkButton(btns, text="Email Back to Author", width=190,
                      fg_color=C_DGOLD,
                      command=self._email_back).pack(side="left", padx=4, pady=8)

    # ── Tree ─────────────────────────────────────────────────────────────────
    def _populate_tree(self):
        self._tree.delete(*self._tree.get_children())
        for c in self.cards:
            vm  = c.get("verification_method","") or "—"
            et  = c.get("equipment_type","") or "—"
            rp  = c.get("responsible_person","") or "—"
            self._tree.insert("","end", values=(
                c["item_num"], c.get("op_no",""),
                (c.get("description","") or "")[:60],
                vm, et, rp))

    # ── Panel ────────────────────────────────────────────────────────────────
    def _flush_panel(self):
        if self._cur is None: return
        c = self.cards[self._cur]
        for k, v in self._tar_vars.items():
            c[k] = v.get()

    def _on_select(self, _=None):
        sel = self._tar_tree_sel()
        if sel is None: return
        self._flush_panel()
        self._cur = sel
        self._populate_tree()
        self._show_panel(self.cards[self._cur])

    def _tar_tree_sel(self):
        sel = self._tree.selection()
        if not sel: return None
        return self._tree.index(sel[0])

    def _show_panel(self, card):
        for w in self._panel.winfo_children():
            w.destroy()
        self._tar_vars = {}

        hf = ctk.CTkFrame(self._panel, fg_color=C_DKGREEN, height=26, corner_radius=0)
        hf.pack(fill="x")
        hf.pack_propagate(False)
        desc_prev = (card.get("description","") or "")[:90]
        ctk.CTkLabel(hf,
                     text=f"  ITEM {card['item_num']}   {desc_prev}",
                     text_color="white",
                     font=ctk.CTkFont("Arial", 9, "bold"),
                     anchor="w").pack(side="left", padx=6, pady=3)

        gf = ctk.CTkFrame(self._panel, fg_color=C_LTGREEN)
        gf.pack(fill="both", expand=True, padx=4, pady=4)
        gf.columnconfigure([1,3], weight=1)

        fields = [
            ("Verification Method:", "verification_method", "dd", VM_OPTIONS),
            ("Equipment Type:",      "equipment_type",       "dd", ET_OPTIONS),
            ("MSA Method:",          "msa_method",           "dd", MSA_OPTIONS),
            ("MSA Ref:",             "msa_ref",               "e",  None),
            ("Responsible Person:",  "responsible_person",    "e",  None),
            ("Control Method:",      "control_method",        "e",  None),
            ("Reaction Plan:",       "reaction_plan",          "e",  None),
            ("Comments:",            "comments",              "e",  None),
        ]
        for i, (lbl, key, wt, opts) in enumerate(fields):
            row, col = divmod(i, 2)
            ctk.CTkLabel(gf, text=lbl, anchor="e",
                         font=ctk.CTkFont("Arial", 9, "bold"),
                         text_color=C_DKGREEN).grid(
                row=row, column=col*2, padx=(6,2), pady=2, sticky="e")
            v = ctk.StringVar(value=card.get(key,""))
            self._tar_vars[key] = v
            if wt == "dd":
                ctk.CTkOptionMenu(gf, variable=v, values=opts, width=175).grid(
                    row=row, column=col*2+1, padx=(0,8), pady=2, sticky="ew")
            else:
                ctk.CTkEntry(gf, textvariable=v).grid(
                    row=row, column=col*2+1, padx=(0,8), pady=2, sticky="ew")

    # ── Holder details ────────────────────────────────────────────────────────
    def _flush_holder(self):
        for k, v in self._h_vars.items():
            self.tar_holder[k] = v.get()

    def _validate_holder(self):
        self._flush_holder()
        missing = [label for label, key in
                   [("TAR Holder Name","name"),("Date","date"),("Employer Number","employer_no")]
                   if not self.tar_holder.get(key,"").strip()]
        return missing

    # ── Actions ───────────────────────────────────────────────────────────────
    def _save_progress(self):
        self._flush_panel()
        self._flush_holder()
        if not self._tip_path:
            self._tip_path = filedialog.asksaveasfilename(
                defaultextension=".tip",
                filetypes=[("TIP Plan","*.tip"),("All Files","*.*")],
                title="Save Plan Progress")
        if not self._tip_path: return
        save_plan(self._tip_path, self.plan, self.cards, self.tar_holder)
        self.app.status(f"Progress saved: {self._tip_path}")

    def _export_approved(self):
        self._flush_panel()
        missing = self._validate_holder()
        if missing:
            messagebox.showwarning("Missing Details",
                "Please complete your details before exporting:\n\n"
                + "\n".join(f"  • {m}" for m in missing))
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".pdf",
            filetypes=[("PDF","*.pdf"),("All Files","*.*")],
            initialfile=f"TIP_{self.plan.get('mva_number','PLAN')}_Approved.pdf",
            title="Save Approved PDF")
        if not path: return
        try:
            export_pdf(path, self.plan, self.cards, self.tar_holder)
            if self._tip_path:
                save_plan(self._tip_path, self.plan, self.cards, self.tar_holder)
            self.app.status(f"Approved PDF saved: {path}")
            if messagebox.askyesno("PDF Saved","Open the PDF now?"):
                os.startfile(path)
        except Exception as ex:
            messagebox.showerror("PDF Error", str(ex))

    def _email_back(self):
        self._flush_panel()
        missing = self._validate_holder()
        if missing:
            messagebox.showwarning("Missing Details",
                "Please complete your details before emailing:\n\n"
                + "\n".join(f"  • {m}" for m in missing))
            return

        author_email = self.plan.get("tar_email","")
        if not author_email:
            dlg = ctk.CTkInputDialog(
                text="Enter the Author's email address:", title="Email Back to Author")
            author_email = dlg.get_input()
        if not author_email: return

        signed_pdf = filedialog.askopenfilename(
            filetypes=[("PDF","*.pdf"),("All Files","*.*")],
            title="Select the Signed PDF to Attach")
        if not signed_pdf: return

        p  = self.plan
        th = self.tar_holder
        body = (
            f"Dear {p.get('mva_author','Author')},\n\n"
            f"Please find attached the approved and signed Test & Inspection Plan.\n\n"
            f"  MVA Number    : {p.get('mva_number','')}\n"
            f"  Approved By   : {th.get('name','')}\n"
            f"  Employer No.  : {th.get('employer_no','')}\n"
            f"  Approval Date : {th.get('date','')}\n\n"
            f"Kind regards,\n{th.get('name','[TAR Holder]')}"
        )
        subject = f"APPROVED: Test & Inspection Plan - {p.get('mva_number','')}"

        if OUTLOOK_OK:
            try:
                send_outlook(author_email, subject, body, [signed_pdf])
                self.app.status("Email draft opened in Outlook.")
            except Exception as ex:
                messagebox.showerror("Email Error", str(ex))
        else:
            messagebox.showinfo("Email (Outlook not available)",
                f"To: {author_email}\nSubject: {subject}\n\nSigned PDF: {signed_pdf}")


# ── Entry point ───────────────────────────────────────────────────────────────
if __name__ == "__main__":
    app = TIPApp()
    app.mainloop()
